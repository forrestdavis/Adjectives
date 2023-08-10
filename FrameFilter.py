import pandas as pd
import openpyxl

class2adj = {'TOUGH': ['bad', 'evil', 'good', 'hard', 'impossible'], 
             'TASTY': ['basic', 'fun', 'spook', 'soft', 'yum'],
             'TALL': ['big', 'chill', 'little', 'long', 'rich'], 
             'PRETTY': ['beautiful', 'love', 'posh', 'pretty', 'splendid'],
             'COLOR': ['blonde', 'blue', 'red', 'rose', 'green'], 
             'SMART': ['cruel', 'fool', 'nice', 'silly', 'mean'], 
             'EMOTION': ['crab', 'happy', 'thank', 'afraid',
                         'sad'],
            }
INF_RIGHT = {'v', 'adv', 'part', 'aux', 'co', 'mod', 'neg', 'cop'}

def ChunkUtterances(data):

    modifiers = {'qn', 'det', 'adj'}
    heads = {'n', 'pro'}

    tagged = data['part_of_speech'].tolist()
    stems = data['stem'].tolist()

    NEWTAGS = []
    NEWSTEMS = []
    # Iterate over tagged sentences
    for i in range(len(tagged)):
        tags = tagged[i].split(' ')
        stem = stems[i].split(' ')

        # Iterate over tags
        j = 0
        ranges = []
        while j < len(tags):
            tag = tags[j].split(':')[0]
            if ':' in tags[j]:
                info = tags[j].split(':')[1]
            else:
                info = ''
            phrase = None
            start = j
            end = j
            # is a head
            if tag in heads:
                phrase = [tag]
                extra = [info]
                # What comes before (e.g., determiners)
                k = j-1
                while k > -1:
                    pre_tag = tags[k].split(':')[0]
                    if pre_tag not in modifiers:
                        k = -1
                        continue
                    phrase = [pre_tag]+phrase
                    start = k
                    k -= 1
                # What comes after (e.g., compounds)
                m = j+1
                while m < len(tags):
                    next_tag = tags[m].split(':')[0]
                    if next_tag not in {'n'}:
                        m = len(tags)
                        continue
                    if ':' in tags[m]:
                        extra = extra + [tags[m].split(':')[1]]
                    else:
                        extra = extra + ['']
                    phrase = phrase + [next_tag]
                    end = m
                    j += 1
                    m += 1

                if len(phrase) == 1:
                    phraseType = 'BARE'
                elif 'adj' in phrase:
                    phraseType = 'AP'
                else:
                    phraseType = 'NP'

                if 'gerund' in extra:
                    phraseType += ':hasGerund'
                # Gather the start, end, and label
                ranges.append(((start, end+1), phraseType))
            j += 1

        # Now that we have ranges and labels, let's chunk up the stem and tag
        newTags = []
        newStem = []
        # For each tag, if it's in a range grab that label
        # otherwise use the original pos
        k = 0
        while k < len(tags):
            if ranges:
                start = ranges[0][0][0]
                end = ranges[0][0][1]
                if k in range(start, end):
                    label = ranges[0][1]
                    words = '-'.join(stem[start:end])
                    # Ensures we only do this once per chunk
                    k = end-1
                    ranges.pop(0)
                else:
                    label = tags[k]
                    words = stem[k]
            else:
                label = tags[k]
                words = stem[k]
            newTags.append(label)
            newStem.append(words)
            k += 1

        newTags = ' '.join(newTags)
        newStem = ' '.join(newStem)
        NEWTAGS.append(newTags)
        NEWSTEMS.append(newStem)

    data['chunked_tags'] = NEWTAGS
    data['chunked_stem'] = NEWSTEMS
    return data

def ExtractFrame(idx, tagged, stem, pairs):

    entries = []

    tag = tagged[idx].split(':')[0]
    if ':' in tagged[idx]:
        extra = tagged[idx].split(':')[1]
    else:
        extra = ''
    text = stem[idx]

    left = 'NA'
    self = tag
    right = 'NA'
    stacked = 'no'

    # too ADJ and as ADJ as
    if idx-1 > 0:
        if stem[idx-1] == 'too':
            left = 'too'
        elif stem[idx-1] == 'as':
            if idx+1 < len(tagged) and stem[idx+1] == 'as':
                left = 'as'
                right = 'as'

    # ADJ at|because|when|inf|of|for|than
    if idx+2 < len(tagged):
        if stem[idx+1] in {'at', 'because', 'when'}:
            right = stem[idx+1]
        elif tagged[idx+1] == 'inf' and tagged[idx+2] in INF_RIGHT:
            right = 'inf'
        elif stem[idx+1] in {'of', 'for', 'to', 'with'}:
            #right = f"{stem[idx+1]} {tagged[idx+2]}"
            right = f"{stem[idx+1]} XP"
        elif stem[idx+1] == 'than':
            right = "than"

    # stacked adjectives
    if idx+1 < len(tagged):
        if tag == 'adj' and tagged[idx+1] == 'adj':
            stacked = 'yes'
        if idx-1 >= 0:
            if tagged[idx-1] == 'adj' and tag == 'adj':
                stacked = 'yes'

    # ADJ gerund
    if extra == 'hasGerund':
        right = 'gerund'
    if idx+1 < len(tagged):
        if 'hasGerund' in tagged[idx+1]:
            right = 'gerund'
    if idx+2 < len(tagged):
        if tagged[idx+1] == 'part':
            if tagged[idx+2] != 'v':
                right = 'gerund'

    #Grab the adjective label
    if tag == 'AP':
        count = 0
        sub_entries = []
        for i in range(len(text.split('-'))):
            if pairs[i][1] == 'adj':
                count += 1
                sub_entries.append([pairs[i][0], left, self, right])
        if count > 1:
            stacked = 'yes'
        for j in range(len(sub_entries)):
            sub_entries[j] = tuple(sub_entries[j]+[stacked])
        entries.extend(sub_entries)
    else:
        entries.append((text, left, self, right, stacked))

    return entries

def GatherUtterances(outname, fname='AllChildesAdjectiveUtterances.tsv', 
                    #tier=['Child', 'Target_Child'], 
                     tier = ['Mother'],
                    StrictCheck=False):

    data = pd.read_csv(fname, sep='\t')

    # Restrict data to target type (e.g., Mother)
    data = data[data['speaker_role'].isin(tier)]

    # First let's add columns where we chunk NPs
    data = ChunkUtterances(data)

    newData = {}
    header = data.columns.tolist()
    for head in header:
        newData[head] = []
    newData['adjective']  = []
    newData['left'] = []
    newData['self'] = []
    newData['right'] = []
    newData['stacked'] = []

    for idx, row in data.iterrows():

        # get word, pos labels
        words = row['stem'].split(' ')
        POS = row['part_of_speech'].split(' ')
        pairs = list(zip(words, POS))

        tagged = row['chunked_tags'].split(' ')
        stem = row['chunked_stem'].split(' ')
        for idx, (tag, text) in enumerate(zip(tagged,stem)):
            tag = tag.split(':')[0]
            size = len(text.split('-'))
            if tag in {'adj', 'AP'}:
                # Get frames
                frames = ExtractFrame(idx, tagged, stem, pairs)
                for frame in frames:
                    for head in header:
                        newData[head].append(row[head])
                    newData['adjective'].append(frame[0])
                    newData['left'].append(frame[1])
                    newData['self'].append(frame[2])
                    newData['right'].append(frame[3])
                    newData['stacked'].append(frame[4])

            pairs = pairs[size:]

    data = pd.DataFrame.from_dict(newData)
    data.to_csv('ChildesFullAdjectivesWithFrames.tsv', sep='\t', index=False)

    adjectives = []

    for adj in class2adj.values():
        adjectives.extend(adj)

    data = data[data['adjective'].isin(adjectives)]

    TYPE = []
    adj2class = {}
    for c in class2adj:
        for adj in class2adj[c]:
            adj2class[adj] = c

    for a in data['adjective']:
        TYPE.append(adj2class[a])

    data['adjective_class'] = TYPE

    for adj in adj2class:
        if StrictCheck:
            assert adj in data['adjective'].unique(), f"Missing {adj}"
        else:
            if adj not in data['adjective'].unique():
                print(f"Missing {adj}")

    mod = []
    for adj in data['adjective']:
        if adj == 'crab':
            mod.append('crabby')
        elif adj == 'love':
            mod.append('lovely')
        elif adj == 'fun':
            mod.append('funny')
        elif adj == 'rose':
            mod.append('rosy')
        elif adj == 'spook':
            mod.append('spooky')
        elif adj == 'yum':
            mod.append('yummy')
        elif adj == 'chill':
            mod.append('chilly')
        elif adj == 'fool':
            mod.append('foolish')
        else:
            mod.append(adj)

    data['adjective'] = mod
    data.target_child_age = data.target_child_age.fillna(-1)
    data.target_child_id = data.target_child_id.fillna(-1)
    data = data.fillna('')

    data.to_csv(outname, sep='\t', index=False)

def toCount(fname):

    data = pd.read_csv(fname, sep='\t', keep_default_na=False)
    TARGETS = ['inf', 'gerund', 'for XP', 
               'to XP', 'with XP', 'of XP']
    data = data[data['right'].isin(TARGETS)]

    COUNTS = {'frame': ['inf', 'gerund', 'for XP', 
                        'to XP', 'with XP', 'of XP']}
    COUNTS['TOUGH'] = [0]*len(COUNTS['frame'])
    COUNTS['SMART'] = [0]*len(COUNTS['frame'])
    COUNTS['TALL'] = [0]*len(COUNTS['frame'])
    COUNTS['TASTY'] = [0]*len(COUNTS['frame'])
    COUNTS['PRETTY'] = [0]*len(COUNTS['frame'])
    COUNTS['EMOTION'] = [0]*len(COUNTS['frame'])
    #COUNTS['COLOR'] = [0]*len(COUNTS['frame'])

    # Get counts
    for _, row in data.iterrows():
        frame = row['right']
        adjClass = row['adjective_class']
        if adjClass == 'COLOR':
            continue
        idx = COUNTS['frame'].index(frame)
        COUNTS[adjClass][idx] += 1

    # Normalize and make into entry like XX% (N)
    for adjClass in COUNTS:
        if adjClass == 'frame': 
            continue
        total = sum(COUNTS[adjClass])
        for idx in range(len(COUNTS[adjClass])):
            N = COUNTS[adjClass][idx]
            percent = int(N/total * 100)
            COUNTS[adjClass][idx] = f"{percent}% ({N})"

    counts = pd.DataFrame.from_dict(COUNTS)
    toDocx(counts)

def toDocx(data):

    import docx
    # Initialise the Word document
    doc = docx.Document()
    # Initialise the table
    t = doc.add_table(rows=data.shape[0]+1, cols=data.shape[1])
    # Add borders
    t.style = 'TableGrid'
    # add the header rows.
    for j in range(data.shape[-1]):
        t.cell(0,j).text = data.columns[j]

    # add the rest of the data frame
    for i in range(data.shape[0]):
        for j in range(data.shape[-1]):
            t.cell(i+1,j).text = str(data.values[i,j])

    # Save the Word doc
    doc.save('table.docx')

def toExcel(fname):
    data = pd.read_csv(fname, sep='\t')
    data.to_excel('../'+fname.split('.')[0]+'.xlsx', index=False)

def basicPlot(fname):
    import seaborn as sns
    import matplotlib.pyplot as plt
    data = pd.read_csv(fname, sep='\t')

    data = data[data['right'].notna()]
    data = data[(data['adjective_class'].isin(['TALL', 'TOUGH']))]
    data.left = data.left.fillna('NA')

    summary = data.groupby(['adjective_class', 'left',
                            'right'])['adjective_class'].count()

    subset=data[(data['right'] == 'inf') & (data['adjective_class'] == 'TALL')]
    subset = subset[subset['left'] != 'too']
    subset = subset[["stem", "adjective", "left", "self", "right", "stacked"]]
    print(subset)

    '''
    sns.set_theme(style="darkgrid")
    sns.displot(data, 
                x='adjective_class', 
                col='right')
    plt.show()
    '''

def ForceAlign(base_fname, extend_fname):

    base = pd.read_excel(base_fname, keep_default_na=False)
    extend = pd.read_excel(extend_fname)

    annotated = base[base['complement'] != '']
    base_heading = base.columns.tolist()
    extend_heading = extend.columns.tolist()
    missing = []
    for bh in base_heading: 
        if bh not in extend_heading:
            missing.append(bh)

    for m in missing:
        extend[m] = ''

    newData = {}
    for bh in base_heading:
        newData[bh] = []

    term2comp = {}
    for _, row in extend.iterrows():
        TC = ((annotated['chunked_stem'] == row['chunked_stem']) & 
              (annotated['adjective'] == row['adjective'] ) &
              (annotated['id'] == row['id']))

        megged = annotated[TC]
        if not megged['adjective'].tolist():
            for entry in base_heading:
                newData[entry].append(row[entry])
            continue
        elif len(megged['adjective'].tolist()) > 1:
            for entry in base_heading:
                if entry in {'complement', 'left', 'right', '+2'}:
                    continue
                newData[entry].append(row[entry])

            ids = megged['id'].tolist()[0]
            adj = megged['adjective'].tolist()[0]
            stem = megged['chunked_stem'].tolist()[0]
            term = f"{ids}+{adj}+{stem}"

            if term not in term2comp:
                comps = megged['complement'].tolist()
                left = megged['left'].tolist()
                right = megged['right'].tolist()
                two = megged['+2'].tolist()
                comp = list(zip(comps, left, right, two))
                term2comp[term] = comp

            info = term2comp[term].pop(0)
            assert len(info) == 4, info
            complement, left, right, two = info
            newData['complement'].append(complement)
            newData['left'].append(left)
            newData['right'].append(right)
            newData['+2'].append(two)

        else:
            for entry in base_heading:
                newData[entry].append(megged[entry].tolist()[0])

    newData = pd.DataFrame.from_dict(newData)
    newData.to_excel('AnnotatedAdjectivesWithFrames.xlsx', index=False)

def BroadenLabeling(fname):
    data = pd.read_csv(fname, sep='\t', keep_default_na=False)

    tagged = data['chunked_tags'].tolist()
    stems = data['chunked_stem'].tolist()
    classes = data['adjective_class'].tolist()

    COUNTS = {'frame': ['inf', 'gerund', 'for XP', 
                        'to XP', 'with XP', 'of XP']}
    COUNTS['TOUGH'] = [0]*len(COUNTS['frame'])
    COUNTS['SMART'] = [0]*len(COUNTS['frame'])
    COUNTS['TALL'] = [0]*len(COUNTS['frame'])
    COUNTS['TASTY'] = [0]*len(COUNTS['frame'])
    COUNTS['PRETTY'] = [0]*len(COUNTS['frame'])
    COUNTS['EMOTION'] = [0]*len(COUNTS['frame'])

    for tag, stem, adjClass in zip(tagged, stems, classes):

        if adjClass not in COUNTS:
            continue

        info = {'inf': 0, 
                'gerund': 0, 
                'for XP': 0, 
                'to XP': 0, 
                'with XP': 0, 
                'of XP': 0}

        POSs = tag.split(' ')
        words = stem.split(' ')
        assert len(POSs) == len(words)
        for pos, word in zip(POSs, words):
            frame = ''
            if pos == 'prep':
                if word in {'to', 'of', 'with', 'for'}:
                    frame = f"{word} XP"
            elif 'gerund' in pos or 'part' == pos:
                frame = 'gerund'
            elif pos == 'inf':
                frame = 'inf'

            if frame != '':
                info[frame] = 1

        for frame in info:
            if info[frame] == 1:
                idx = COUNTS['frame'].index(frame)
                COUNTS[adjClass][idx] += 1

    # Normalize and make into entry like XX% (N)
    for adjClass in COUNTS:
        if adjClass == 'frame': 
            continue
        total = sum(COUNTS[adjClass])
        for idx in range(len(COUNTS[adjClass])):
            N = COUNTS[adjClass][idx]
            percent = int(N/total * 100)
            COUNTS[adjClass][idx] = f"{percent}% ({N})"

    counts = pd.DataFrame.from_dict(COUNTS)
    print(counts)
    toDocx(counts)

def toAnnotatedSubset(fname):

    data = pd.read_csv(fname, sep='\t')
    TARGETS = ['inf', 'gerund', 'for XP', 
               'to XP', 'with XP', 'of XP']
    data = data[data['right'].isin(TARGETS)]
    data = data[data['stacked'] == 'no']

    glosses = data['gloss'].tolist()
    adjectives = data['adjective'].tolist()
    
    REPLACEMENTS = { 
                    'good': ['better', 'best'], 
                    'silly': ['silliest'], 
                    'funny': ['fun', 'funniest'], 
                    'bad': ['worse', 'worst'],
                    'happy': ['happier', 'happiest'], 
                    'pretty': ['prettiest'],
                    'long': ['longer', 'longest'], 
                    'big': ['bigger', 'biggest'],
                    'nice': ['nicer', 'nicest'],
                    'hard': ['hardest', 'harder'],
                    'mean': ['meanest'],
                    'soft': ['softest'],
                    'blue': ["blue's"],
                    'green': ["green's"],
                    'thank': ['thankful'],
                   }

    MASKED = []
    for gloss, adj in zip(glosses, adjectives):
        adj = adj + ' '
        if adj not in gloss:
            if ' gud ' in gloss or ' betta ' in gloss or 'gooder' in gloss:
                gloss = 'NA'
            else:
                replacement = REPLACEMENTS[adj.strip()]
                beenReplaced = 0
                for replace in replacement:
                    replace = replace + ' '
                    if replace in gloss:
                        gloss = gloss.replace(replace, adj)
                        beenReplaced = 1
                        break
                assert beenReplaced, f"{adj} {replacement} {gloss}"
        if gloss == 'NA':
            masked = 'NA'
        else:
            assert adj in gloss, f"{adj} not in {gloss}"
            if gloss.count(adj) != 1:
                masked = 'NA'
            else:
                assert gloss.count(adj) == 1, f"too many '{adj}' in '{gloss}'"
                masked = gloss.replace(adj, 'MASKTOKEN ')

        import re
        if re.findall(f'\wMASKTOKEN', masked):
            masked = 'NA'
        MASKED.append(masked)
    data['masked'] = MASKED
    data = data[data['masked'] != 'NA']
    data.to_csv('AnnotatedSubsetForModels.tsv', sep='\t', index=False)

if __name__ == '__main__':

    # Downloads CHILDES data (via R script)

    # Gathers utterances just spoken by Mother tier
    fname = 'MotherRestrictedFrames.tsv'
    #GatherUtterances(fname)

    # Makes Annotated Subset
    #toAnnotatedSubset(fname)

    #toCount(fname)
    #BroadenLabeling(fname)
